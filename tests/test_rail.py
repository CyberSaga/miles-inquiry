"""高鐵表 距離 from a THSR chainage table + rail_distance_overrides.csv (synthetic data; the shipped THSR table is checked too)."""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from miles_inquiry.cli import main
from miles_inquiry.fill import Options, fill_workbook, write_rail_todo
from miles_inquiry.overrides import Overrides
from miles_inquiry.rail import RailData, station_key

# made-up distances, NOT the real THSR numbers
STATIONS = ("\ufefffrom,to,kg_co2e,km\n#c\n台北,板橋,,7\n台北,台中,,160\n板橋,台中,,153\n"
            "台北,左營,,340.5\n台北,彰化,,190\n")
LOOKUP = ("﻿value,station,review,note\n高雄,左營,,高雄站即左營\n粄橋,板橋,yes,疑為板橋誤植\n"
          "斗六,,,斗六為台鐵站\n苏州北,蘇州北,,\n")
OVERRIDES = "﻿from,to,km,source\n#c\n徐州東,蘇州北,300,12306\n"


@pytest.fixture
def rail_dir(tmp_path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    (d / "thsr_distances.csv").write_text(STATIONS, encoding="utf-8")
    (d / "hsr_lookup.csv").write_text(LOOKUP, encoding="utf-8")
    (d / "rail_distance_overrides.csv").write_text(OVERRIDES, encoding="utf-8")
    return d


def test_station_key():
    assert station_key("臺北車站") == "台北" and station_key(" 彰化站 ") == "彰化"
    assert station_key("左營高鐵站") == "左營" and station_key("站") == "站"


def test_distance_rules(rail_dir):
    rail = RailData.load(rail_dir)
    assert rail.distance("台北", "台中")[0].km == 160
    assert rail.distance("臺中", "台北車站")[0].km == 160          # 臺/台, 車站 suffix, either direction
    d, ra, rb = rail.distance("高雄", "台北")
    assert d.km == 340.5 and ra.station == "左營"
    assert rail.distance("苏州北", "徐州東")[0].km == 300            # alias merges spellings for the override
    assert rail.distance("台中", "左營")[0].source == "missing_pair"
    assert rail.distance("台北", "臺北車站")[0].km == 0
    d, _, rb = rail.distance("台北", "斗六")
    assert d.km is None and rb.non_station
    assert rail.distance("南京南", "全椒")[0].km is None


def build(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "航空"
    ws.append(["公司", "航班(起站)", "航班(迄站)", "距離", "人次", "碳排量", "備註"])
    hsr = wb.create_sheet("高鐵")
    hsr.append(["公司", "單位代號", "單位名稱", "日期", "高鐵(起站)", "高鐵(迄站)", "距離", "人次", "碳排量", "備註"])
    rows = [
        ("台中", "桃園", 1, None),      # 2: J2 is the factor cell – nothing may be written into it
        ("臺北", "臺中", 2, None),      # 3
        ("高雄", "台北", 1, None),      # 4: alias
        ("粄橋", "台中", 1, None),      # 5: review alias → note
        ("台北", "斗六", 1, None),      # 6: non-station
        ("苏州北", "徐州東", 1, None),  # 7: override
        ("南京南", "全椒", 2, None),    # 8: missing pair
        ("全椒", "南京南", 1, None),    # 9: same pair, other direction
        ("無", "無", 0, "排除"),        # 10: skipped
        ("台北", "火星", 1, None),      # 11: unknown value
        ("台中", "左營", 1, None),      # 12: both stations, pair missing from the table
        ("台北", "臺北車站", 1, None),  # 13: same station
    ]
    for o, d, pax, note in rows:
        hsr.append(["X", "B", "C", None, o, d, None, pax, None, note])
    hsr["J2"] = 0.029
    hsr.append(["X", "B", "C", None, "台北", "左營", 345, 1, None, None])   # 14: hand-typed distance kept
    wb.save(path)
    return path


def run(path, lookup, rail_dir, out, **kw):
    kw.setdefault("hsr_factor_cell", "J2")      # the synthetic workbooks keep the factor in J2; the shipped default is K2
    return fill_workbook(path, out, lookup, Overrides({}), Options(**kw), RailData.load(rail_dir))


def test_fill_hsr_sheet(lookup, rail_dir, tmp_path):
    out = tmp_path / "o.xlsx"
    s = run(build(tmp_path / "w.xlsx"), lookup, rail_dir, out)
    ws = openpyxl.load_workbook(out)["高鐵"]
    g = {r: ws.cell(r, 7).value for r in range(2, 15)}
    j = {r: ws.cell(r, 10).value for r in range(2, 15)}
    assert g[2] is None and j[2] == 0.029                       # 桃園 not in synthetic table, factor untouched
    assert g[3] == 160 and isinstance(g[3], int) and j[3] is None
    assert g[4] == 340.5 and j[4] is None                        # 高雄→左營 is report-only
    assert g[5] == 153 and j[5] == "起站 粄橋→板橋：疑為板橋誤植"
    assert g[6] is None and j[6] == "迄站「斗六」非高鐵站（斗六為台鐵站），請確認"
    assert g[7] == 300 and j[7] is None
    assert g[8] is None and "清單中未有此起迄組合之距離，請在 data/rail_distance_overrides.csv 填距離" in j[8] and "非台灣高鐵路段" not in j[8]
    assert g[10] is None and j[10] == "排除"
    assert g[11] is None and j[11] == "無法對應高鐵站：火星，請更新 data/hsr_lookup.csv"
    assert g[12] is None and j[12] == "data/thsr_distances.csv 缺 台中-左營 距離" and s.hsr_missing_thsr["台中-左營"] == 1
    assert g[13] == 0 and j[13] == "起迄同站，距離為 0，請確認"
    assert g[14] == 345 and s.hsr_distance_counts["kept"] == 1
    assert s.hsr_distance_counts["thsr"] == 4 and s.hsr_distance_counts["override"] == 1
    assert s.hsr_missing_pairs == {("全椒", "南京南"): [8, 9]}
    # the 碳排量 formula picks the new distance up
    assert ws["I3"].value == '=IF(G3="","",G3*H3*$J$2)'
    text = s.console_summary()
    assert "台中-左營（1 列）" in text and "火星" in text and "缺距離的站對 1 組" in text

    todo = write_rail_todo(s, tmp_path / "todo.csv")
    rows = list(csv.reader(todo.read_text(encoding="utf-8-sig").splitlines()))
    assert rows == [["from", "to", "km", "source", "rows"], ["全椒", "南京南", "", "", "8 9"]]


def test_hsr_rerun_idempotent_and_flag(lookup, rail_dir, tmp_path):
    src = build(tmp_path / "w.xlsx")
    run(src, lookup, rail_dir, tmp_path / "1.xlsx")
    run(tmp_path / "1.xlsx", lookup, rail_dir, tmp_path / "2.xlsx")
    a, b = (openpyxl.load_workbook(tmp_path / n)["高鐵"] for n in ("1.xlsx", "2.xlsx"))
    assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()]
    run(src, lookup, rail_dir, tmp_path / "n.xlsx", hsr_distance=False)
    assert openpyxl.load_workbook(tmp_path / "n.xlsx")["高鐵"]["G3"].value is None


def test_shipped_thsr_table_reproduces_official_footprint():
    """data/thsr_distances.csv: 66 pairs, km × 0.029 rounds back to THSRC's published kg CO2e."""
    path = Path(__file__).resolve().parent.parent / "data" / "thsr_distances.csv"
    rows = [r for r in csv.DictReader(ln for ln in path.read_text(encoding="utf-8-sig").splitlines()
                                       if not ln.startswith("#"))]
    assert len(rows) == 66 and len({r["from"] for r in rows} | {r["to"] for r in rows}) == 12
    for r in rows:
        assert round(float(r["km"]) * 0.029 + 1e-9, 2) == float(r["kg_co2e"]), r
    rail = RailData.load(path.parent)
    assert rail.distance("臺北", "左營")[0].km == 343.4 and rail.distance("台中", "台北")[0].km == 161.7


def _write(d: Path, stations=STATIONS, lookup=LOOKUP, overrides=OVERRIDES) -> Path:
    d.mkdir(exist_ok=True)
    (d / "thsr_distances.csv").write_text(stations, encoding="utf-8")
    (d / "hsr_lookup.csv").write_text(lookup, encoding="utf-8")
    (d / "rail_distance_overrides.csv").write_text(overrides, encoding="utf-8")
    return d


def test_override_rows_merge_spellings_through_lookup(tmp_path):
    lookup = "value,station,review,note\n苏州北,蘇州北,,\n徐州东,徐州東,,\n"
    # typed in either script, either direction; the same pair twice with the same km is fine
    overrides = "from,to,km,source\n苏州北,徐州东,300,a\n徐州東,蘇州北,300,b\n"
    rail = RailData.load(_write(tmp_path / "d1", lookup=lookup, overrides=overrides))
    assert len(rail.overrides) == 1
    for a, b in (("苏州北", "徐州东"), ("蘇州北", "徐州東"), ("徐州东", "蘇州北")):
        assert rail.distance(a, b)[0].km == 300
    clash = "from,to,km,source\n苏州北,徐州东,300,a\n徐州東,蘇州北,301,b\n"
    with pytest.raises(ValueError, match="出現兩次且距離不同"):
        RailData.load(_write(tmp_path / "d2", lookup=lookup, overrides=clash))
    # a fifth column (e.g. the worksheet rows of the todo file) is ignored, never stored
    rail = RailData.load(_write(tmp_path / "d3", lookup=lookup, overrides="from,to,km,source,rows\n苏州北,徐州东,300,a,12 45\n"))
    assert rail.overrides == {("徐州東", "蘇州北"): (300.0, "a")}


def test_is_non_thsr_only_for_listed_pairs(tmp_path):
    overrides = OVERRIDES + "台北,台中,150,pinned Taiwan pair\n台北,蘇州北,999,mixed pair\n"
    rail = RailData.load(_write(tmp_path / "d", overrides=overrides))
    assert rail.is_non_thsr("苏州北", "徐州東") and rail.is_non_thsr("徐州東", "蘇州北")   # listed, either spelling/direction
    assert rail.is_non_thsr("台北", "苏州北")                    # listed, one end is a THSR station
    assert not rail.is_non_thsr("台北", "台中")                  # a pinned pair of two THSR stations stays Taiwan
    assert not rail.is_non_thsr("南京南", "全椒")                # China, but not listed (yet): no guessing
    assert not rail.is_non_thsr("火星", "月球")                  # junk / mistyped Taiwan stations are never "China"
    assert not rail.is_non_thsr("台北", "斗六") and not rail.is_non_thsr("斗六", "蘇州北")   # known non-station
    assert rail.distance("台北", "台中")[0].km == 150            # the pinned distance itself still applies


def build_china(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "航空"
    wb.active.append(["公司", "航班(起站)", "航班(迄站)", "距離", "人次", "碳排量", "備註"])
    hsr = wb.create_sheet("高鐵")
    hsr.append(["公司", "單位代號", "單位名稱", "日期", "高鐵(起站)", "高鐵(迄站)", "距離", "人次", "碳排量", "備註", "說明", "中國"])
    for o, d, pax in (("台中", "桃園", 1),      # 2: Taiwan (桃園 missing from the synthetic table)
                      ("台北", "台中", 2),      # 3: Taiwan
                      ("苏州北", "徐州東", 3),  # 4: listed China pair
                      ("南京南", "全椒", 1),    # 5: China, not listed → pending, gets a note
                      ("台北", "斗六", 1),      # 6: known non-station
                      ("徐州東", "苏州北", 1),  # 7: listed China pair, hand-typed number in 碳排量
                      ("苏州北", "徐州東", 1)):  # 8: listed China pair, someone else's formula in 碳排量
        hsr.append(["X", "B", "C", None, o, d, None, pax, None, None])
    hsr["J2"] = 0.029
    hsr["L2"] = 0.0262
    hsr["I7"] = 5.5
    hsr["I8"] = "=G8*H8*0.5"
    wb.save(path)
    return path


F_J = '=IF(G{r}="","",G{r}*H{r}*$J$2)'
F_L = '=IF(G{r}="","",G{r}*H{r}*$L$2)'


def test_china_factor_cell_only_for_listed_pairs(lookup, rail_dir, tmp_path):
    out = tmp_path / "c.xlsx"
    s = run(build_china(tmp_path / "w.xlsx"), lookup, rail_dir, out, hsr_china_factor_cell="L2")
    ws = openpyxl.load_workbook(out)["高鐵"]
    assert ws["I3"].value == F_J.format(r=3)                     # Taiwan keeps J2
    assert ws["I4"].value == F_L.format(r=4)                     # listed China pair
    assert ws["I5"].value == F_J.format(r=5)                     # not listed → default factor, still pending
    assert ws["I6"].value == F_J.format(r=6)                     # known non-station
    assert ws["I7"].value == 5.5 and ws["I8"].value == "=G8*H8*0.5"   # hand-typed values are kept
    assert ws["G4"].value == 300 and ws["L2"].value == 0.0262 and ws["J2"].value == 0.029
    assert s.hsr_carbon_china_rows == 1 and s.hsr_china_factor.value == 0.0262
    assert "非台灣高鐵路段 1 列用係數 0.0262" in s.console_summary()
    assert "注意：" not in s.console_summary()

    # --overwrite replaces hand-typed carbon on the China rows by the China formula, not the Taiwan one
    s = run(build_china(tmp_path / "w1.xlsx"), lookup, rail_dir, tmp_path / "ow.xlsx", hsr_china_factor_cell="L2", overwrite=True)
    ws = openpyxl.load_workbook(tmp_path / "ow.xlsx")["高鐵"]
    assert ws["I7"].value == F_L.format(r=7) and ws["I8"].value == F_L.format(r=8) and s.hsr_carbon_china_rows == 3

    # re-running the output with the same option changes nothing
    run(out, lookup, rail_dir, tmp_path / "e.xlsx", hsr_china_factor_cell="L2")
    a, b = (openpyxl.load_workbook(tmp_path / n)["高鐵"] for n in ("c.xlsx", "e.xlsx"))
    assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()]


def test_without_china_factor_every_row_uses_j2_and_says_so(lookup, rail_dir, tmp_path):
    s = run(build_china(tmp_path / "w.xlsx"), lookup, rail_dir, tmp_path / "d.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "d.xlsx")["高鐵"]
    assert [ws.cell(r, 9).value for r in (2, 3, 4, 5, 6)] == [F_J.format(r=r) for r in (2, 3, 4, 5, 6)]
    assert s.hsr_carbon_china_default == 1                       # row 4 is a listed China pair
    text = s.console_summary()
    assert "注意：1 列是非台灣高鐵路段" in text and "--hsr-china-factor-cell" in text
    # dropping the option on a re-run refreshes our own L2 formulas back to J2 – and the notice says so
    run(build_china(tmp_path / "w2.xlsx"), lookup, rail_dir, tmp_path / "o.xlsx", hsr_china_factor_cell="L2")
    s = run(tmp_path / "o.xlsx", lookup, rail_dir, tmp_path / "o2.xlsx")
    assert openpyxl.load_workbook(tmp_path / "o2.xlsx")["高鐵"]["I4"].value == F_J.format(r=4)
    assert "注意：1 列是非台灣高鐵路段" in s.console_summary()


def test_china_factor_cell_must_be_a_number(lookup, rail_dir, tmp_path):
    first = tmp_path / "first.xlsx"
    run(build_china(tmp_path / "w.xlsx"), lookup, rail_dir, first, hsr_china_factor_cell="L2")
    wb = openpyxl.load_workbook(first)
    wb["高鐵"]["L2"] = "=0.0262"                                  # a formula is not a number for us
    wb.save(first)
    s = run(first, lookup, rail_dir, tmp_path / "o.xlsx", hsr_china_factor_cell="L2")
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx")["高鐵"]
    assert ws["I4"].value is None                                 # the earlier $L$2 formula is cleared, not left stale
    assert ws["I3"].value == F_J.format(r=3) and ws["I5"].value == F_J.format(r=5)
    assert any("L2 不是數字" in w and "不會沿用台灣係數" in w for w in s.warnings)
    # fresh file, blank factor cell: nothing is written for the China rows, no fallback to the Taiwan factor
    src = build_china(tmp_path / "w3.xlsx")
    wb = openpyxl.load_workbook(src)
    wb["高鐵"]["L2"] = None
    wb.save(src)
    run(src, lookup, rail_dir, tmp_path / "b.xlsx", hsr_china_factor_cell="L2")
    assert openpyxl.load_workbook(tmp_path / "b.xlsx")["高鐵"]["I4"].value is None


def test_china_factor_cell_inside_a_written_column_is_refused(lookup, rail_dir, tmp_path):
    s = run(build_china(tmp_path / "w.xlsx"), lookup, rail_dir, tmp_path / "o.xlsx", hsr_china_factor_cell="I9")
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx")["高鐵"]
    assert ws["I4"].value == F_J.format(r=4) and ws["I9"].value is None
    assert any("I9" in w and "會被填入的欄位" in w for w in s.warnings)


@pytest.mark.parametrize("bad", ["L", "2", "foo", "L2:M3", "L0", "ZZZZ99", ""])
def test_cell_options_are_validated(bad):
    with pytest.raises(ValueError, match="不是有效的儲存格位址"):
        Options(hsr_factor_cell=bad)
    if bad:
        with pytest.raises(ValueError, match="不是有效的儲存格位址"):
            Options(hsr_china_factor_cell=bad)
    assert Options(hsr_factor_cell="$j$2", hsr_china_factor_cell="l2").hsr_china_factor_cell == "L2"


def test_china_factor_cell_never_receives_notes(lookup, rail_dir, tmp_path):
    src = build_china(tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(src)
    wb["高鐵"]["L2"] = None
    wb["高鐵"]["J5"] = 0.0262          # parked in the 備註 cell of row 5, which would otherwise receive a note
    wb.save(src)
    run(src, lookup, rail_dir, tmp_path / "o.xlsx", hsr_china_factor_cell="J5")
    ws = openpyxl.load_workbook(tmp_path / "o.xlsx")["高鐵"]
    assert ws["J5"].value == 0.0262 and ws["I4"].value == '=IF(G4="","",G4*H4*$J$5)'


def test_cli_forwards_china_factor_cell(rail_dir, lookup_path, tmp_path, capsys):
    from tests.conftest import SYNTH_AIRPORTS
    ap = tmp_path / "airports.csv"
    with ap.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["iata", "name", "municipality", "iso_country", "type", "lat", "lon", "scheduled_service"])
        for a in SYNTH_AIRPORTS:
            w.writerow([a.iata, a.name, a.municipality, a.iso_country, a.type, a.lat, a.lon, "yes" if a.scheduled_service else "no"])
    src = build_china(tmp_path / "w.xlsx")
    base = ["fill", str(src), "--lookup", str(lookup_path), "--airports", str(ap), "--data-dir", str(rail_dir), "--no-report",
            "--hsr-factor-cell", "J2"]

    assert main(base + ["-o", str(tmp_path / "a.xlsx"), "--hsr-china-factor-cell", "L2"]) == 0
    ws = openpyxl.load_workbook(tmp_path / "a.xlsx")["高鐵"]
    assert ws["I4"].value == F_L.format(r=4) and ws["I3"].value == F_J.format(r=3)

    assert main(base + ["-o", str(tmp_path / "b.xlsx")]) == 0                      # flag absent → all J2
    assert openpyxl.load_workbook(tmp_path / "b.xlsx")["高鐵"]["I4"].value == F_J.format(r=4)
    assert "注意：1 列是非台灣高鐵路段" in capsys.readouterr().out

    # distance filling off, but the China factor still needs the station data
    assert main(base + ["-o", str(tmp_path / "c.xlsx"), "--no-hsr-distance", "--hsr-china-factor-cell", "L2"]) == 0
    ws = openpyxl.load_workbook(tmp_path / "c.xlsx")["高鐵"]
    assert ws["I4"].value == F_L.format(r=4) and ws["G4"].value is None

    # a malformed cell is a normal error message, not a traceback
    assert main(base + ["-o", str(tmp_path / "d.xlsx"), "--hsr-china-factor-cell", "foo"]) == 1
    assert "不是有效的儲存格位址" in capsys.readouterr().err

    # station table missing: the option is reported, not silently ignored
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main(base[:-5] + ["--data-dir", str(empty), "--no-report", "--hsr-factor-cell", "J2", "-o", str(tmp_path / "e.xlsx"), "--hsr-china-factor-cell", "L2"]) == 0
    assert "沒有高鐵站名資料" in capsys.readouterr().out


def test_shipped_rail_data_is_consistent():
    """data/hsr_lookup.csv + rail_distance_overrides.csv: loadable, one canonical spelling per station, no alias chains."""
    data = Path(__file__).resolve().parent.parent / "data"
    rail = RailData.load(data)
    for key, (target, _note, _review) in rail.lookup.items():
        assert not target or target == key or target not in rail.lookup or rail.lookup[target][0] == target, key
    rows = [r for r in csv.DictReader(ln for ln in (data / "rail_distance_overrides.csv").read_text(encoding="utf-8-sig")
                                       .splitlines() if not ln.startswith("#"))]
    assert len(rows) == len(rail.overrides)                           # no pair listed twice
    for r in rows:
        assert float(r["km"]) > 0 and r["source"], r
        for name in (r["from"], r["to"]):
            assert rail.resolve(name).canonical == station_key(name), f"{name} 不是標準寫法"


def _painted(cell) -> bool:
    return (cell.fill.fill_type == "solid" and cell.fill.start_color.rgb == "FFFFFF00"
            and cell.font.color is not None and cell.font.color.rgb == "FFFF0000")


def test_hsr_rows_with_a_request_for_a_human_are_painted(lookup, rail_dir, tmp_path):
    out = tmp_path / "o.xlsx"
    s = run(build(tmp_path / "w.xlsx"), lookup, rail_dir, out)
    ws = openpyxl.load_workbook(out)["高鐵"]
    asks = {5, 6, 8, 9, 11, 12, 13}     # review alias / 非高鐵站 / 缺距離 ×2 / 不認得 / 站對缺 / 起迄同站
    for r in range(3, 15):
        cells = [ws.cell(r, c) for c in (5, 6, 10)]    # 起站、迄站、備註
        assert all(_painted(c) for c in cells) == (r in asks), r
        assert all(_painted(c) for c in cells) or not any(_painted(c) for c in cells), r
        assert not _painted(ws.cell(r, 7)), r
    assert s.flagged_hsr == len(asks) + 1 and "高鐵 8 列" in s.console_summary()   # + row 2 (桃園)
    # a resolved row stays plain: aliases, overrides and hand-typed distances need no human
    assert not any(_painted(ws.cell(r, c)) for r in (3, 4, 7, 14) for c in (5, 6, 10))


def test_hsr_painting_is_idempotent_and_optional(lookup, rail_dir, tmp_path):
    src = build(tmp_path / "w.xlsx")
    run(src, lookup, rail_dir, tmp_path / "1.xlsx")
    run(tmp_path / "1.xlsx", lookup, rail_dir, tmp_path / "2.xlsx")
    for name in ("1.xlsx", "2.xlsx"):
        ws = openpyxl.load_workbook(tmp_path / name)["高鐵"]
        assert all(_painted(ws.cell(r, c)) for r in (6, 8, 11) for c in (5, 6, 10)), name
    run(src, lookup, rail_dir, tmp_path / "n.xlsx", highlight=False)
    ws = openpyxl.load_workbook(tmp_path / "n.xlsx")["高鐵"]
    assert not any(_painted(c) for row in ws.iter_rows() for c in row)


def test_painting_never_touches_the_factor_cell(lookup, rail_dir, tmp_path):
    out = tmp_path / "o.xlsx"
    run(build(tmp_path / "w.xlsx"), lookup, rail_dir, out)
    ws = openpyxl.load_workbook(out)["高鐵"]
    # row 2 (桃園 unknown): its 備註 is the J2 factor cell → only 起站/迄站 are painted, J2 is left alone
    assert _painted(ws["E2"]) and _painted(ws["F2"]) and not _painted(ws["J2"])


def build_layout(path: Path, china=True) -> Path:
    """The bank's layout: Taiwan factor in K2, China factor in K3, sources in L2:L3 (no factor in 備註 column J)."""
    wb = openpyxl.Workbook()
    wb.active.title = "航空"
    wb.active.append(["公司", "航班(起站)", "航班(迄站)", "距離", "人次", "碳排量", "備註"])
    hsr = wb.create_sheet("高鐵")
    hsr.append(["公司", "單位代號", "單位名稱", "日期", "高鐵(起站)", "高鐵(迄站)", "距離", "人次", "碳排量", "備註"])
    for o, d in (("台北", "台中"), ("苏州北", "徐州東"), ("南京南", "全椒")):
        hsr.append(["X", "B", "C", None, o, d, None, 1, None, None])
    hsr["K2"], hsr["L2"] = 0.029, "環境部產品碳足跡標籤-2025-高速鐵路旅客運輸服務-高鐵"
    if china:
        hsr["K3"], hsr["L3"] = 0.0262, "中国产品全生命周期温室气体排放系数-2022-高铁"
    wb.save(path)
    return path


def test_default_layout_taiwan_k2_china_k3(lookup, rail_dir, tmp_path):
    out = tmp_path / "o.xlsx"
    s = fill_workbook(build_layout(tmp_path / "w.xlsx"), out, lookup, Overrides({}), Options(), RailData.load(rail_dir))
    ws = openpyxl.load_workbook(out)["高鐵"]
    assert ws["I2"].value == '=IF(G2="","",G2*H2*$K$2)'                    # Taiwan
    assert ws["I3"].value == '=IF(G3="","",G3*H3*$K$3)'                    # listed China pair
    assert ws["I4"].value == '=IF(G4="","",G4*H4*$K$2)'                    # unlisted China pair: pending, default factor
    assert ws["J4"].value == "清單中未有此起迄組合之距離，請在 data/rail_distance_overrides.csv 填距離"
    assert ws["K2"].value == 0.029 and ws["K3"].value == 0.0262 and ws["L3"].value.startswith("中国产品")
    assert s.hsr_china_factor.value == 0.0262 and s.hsr_carbon_china_rows == 1
    assert not any("高鐵" in w for w in s.warnings)
    # nothing changes on a re-run, with no flag at all
    fill_workbook(out, tmp_path / "o2.xlsx", lookup, Overrides({}), Options(), RailData.load(rail_dir))
    a, b = (openpyxl.load_workbook(tmp_path / n)["高鐵"] for n in ("o.xlsx", "o2.xlsx"))
    assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()]


def test_default_china_cell_is_optional(lookup, rail_dir, tmp_path):
    out = tmp_path / "o.xlsx"
    s = fill_workbook(build_layout(tmp_path / "w.xlsx", china=False), out, lookup, Overrides({}), Options(),
                      RailData.load(rail_dir))
    ws = openpyxl.load_workbook(out)["高鐵"]
    assert ws["I3"].value == '=IF(G3="","",G3*H3*$K$2)'                    # no K3 → one factor, no warning
    assert not any("高鐵" in w for w in s.warnings) and s.hsr_china_factor is None and s.hsr_carbon_china_default == 1
    assert "請填在 K3" in s.console_summary()


def test_explicit_china_cell_overrides_the_k3_default(lookup, rail_dir, tmp_path):
    src = build_layout(tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(src)
    wb["高鐵"]["M2"] = 0.5
    wb.save(src)
    out = tmp_path / "o.xlsx"
    fill_workbook(src, out, lookup, Overrides({}), Options(hsr_china_factor_cell="M2"), RailData.load(rail_dir))
    assert openpyxl.load_workbook(out)["高鐵"]["I3"].value == '=IF(G3="","",G3*H3*$M$2)'
