"""End-to-end tests on a SYNTHETIC workbook built here (never the real sample)."""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from miles_inquiry.cli import main
from miles_inquiry.fill import Options, append_note, default_output_path, fill_workbook, find_header, write_report
from miles_inquiry.overrides import Overrides

HEADERS = ["公司", "單位代號", "單位名稱", "日期", "航班(起站)", "航班(迄站)", "距離", "人次", "碳排量", "備註"]

# (origin, dest, pax, existing_note, existing_distance)
ROWS = [
    ("SHA", "TSA", 1, None, None),                     # 2: direct codes (SHA not in synth DB → unresolved)
    ("無", "無", 0, "排除：人次為0或起迄為「無」", None),  # 3: skipped
    ("澎湖", "高雄", 2, None, None),                   # 4: lookup names
    ("新竹", "台北", 1, None, None),                   # 5: HSR
    ("台中", "台北", 1, None, None),                   # 6: HSR
    ("TAS", "MZG", 1, None, None),                     # 7: correction
    ("MZG", "kHH", 1, "既有備註", None),               # 8: lowercase + existing note
    ("TPE", "JFK", 3, None, None),                     # 9: long haul
    ("TPE", "HONG KONG", 1, None, None),               # 10: name with spaces
    ("火星", "TPE", 1, None, None),                    # 11: unresolved
    ("TPE", "KHH", 1, None, 999),                      # 12: HSR-like with existing distance
    ("MZG", "TSA", 1, None, 123),                      # 13: existing distance kept
    ("0", "TPE", 1, None, None),                       # 14: skipped by '0'
    ("TPE", "MZG", 0, None, None),                     # 15: skipped by pax 0
    ("TPE", "SIA", 1, None, None),                     # 16: correction SIA→SAI
    ("MZG", "MZG", 1, None, None),                     # 17: same airport
    ("台北", "澎湖", 1, None, None),                   # 18: review-flagged lookup (台北→TSA) → note in 備註
    ("TPE", "PNH", 1, None, None),                     # 19: direct code without scheduled service
]


def build_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "航空"
    ws.append(HEADERS)
    ws["A1"].font = Font(name="微軟正黑體", bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor="FFDDDDDD")
    ws.column_dimensions["E"].width = 15.9
    for i, (o, d, pax, note, dist) in enumerate(ROWS, start=2):
        ws.append(["測試公司", f"A{i:04d}", "測試分行", dt.datetime(2026, 5, 1), o, d, dist, pax, None, note])
    ws["I2"].value = "=距離*人次*係數"           # literal TEXT that looks like a formula (as in real workbooks)
    ws["I2"].data_type = "s"
    ws["I3"] = "=G3*H3*0.123"                    # real formula
    ws["L2"] = "https://www.greatcirclemap.com/"  # ad-hoc reference note (column L)
    ws["K5"] = "誤植，交通方式應調整為高鐵"
    other = wb.create_sheet("高鐵")
    other.append(["公司", "高鐵(起站)", "高鐵(迄站)", "距離", "人次"])
    other.append(["測試公司", "台北", "左營", 345, 2])
    other["F2"] = "=D2*E2"
    wb.save(path)
    return path


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def workbook(tmp_path) -> Path:
    return build_workbook(tmp_path / "synthetic.xlsx")


def run(workbook: Path, lookup, out: Path, **kw):
    return fill_workbook(workbook, out, lookup, Overrides({}), Options(**kw))


def test_find_header_tolerates_whitespace_and_fullwidth(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "航空"
    ws.append(["title row"])
    ws.append(["公司", " 航班（起站） ", "航班(迄站)\n", "距離 ", "人次", "備註"])
    p = tmp_path / "h.xlsx"
    wb.save(p)
    row, cols = find_header(openpyxl.load_workbook(p)["航空"])
    assert row == 2 and cols == {"origin": 2, "dest": 3, "distance": 4, "pax": 5, "note": 6}


def test_append_note_idempotent():
    assert append_note(None, "a") == "a"
    assert append_note("x", "a") == "x；a"
    assert append_note("x；a", "a") == "x；a"
    assert append_note("x; a ", "a") == "x; a "
    assert append_note("", "a") == "a"
    # notes that themselves contain the separator must not be re-appended
    assert append_note("x", "a；b") == "x；a；b"
    assert append_note("x；a；b", "a；b") == "x；a；b"
    assert append_note("a；x；b", "a；b") == "a；x；b；a；b"
    assert append_note("x", "") == "x"


def test_end_to_end(workbook, lookup, tmp_path):
    before = sha(workbook)
    out = tmp_path / "out.xlsx"
    summary = run(workbook, lookup, out)
    assert sha(workbook) == before, "input file must not be modified"
    assert out.exists()

    by_row = {r.row: r for r in summary.results}
    assert summary.counts["skipped"] == 3 and {3, 14, 15} == {r for r, x in by_row.items() if x.status == "skipped"}
    assert by_row[4].status == "filled" and by_row[4].km == pytest.approx(132, abs=1)
    assert by_row[4].origin_iata == "MZG" and by_row[4].dest_iata == "KHH"
    assert by_row[4].url == "https://www.greatcirclemap.com/?routes=MZG-KHH"
    assert by_row[5].status == "hsr" and by_row[5].km is None
    assert by_row[6].status == "hsr" and by_row[6].km is None
    assert by_row[7].status == "filled" and by_row[7].origin_iata == "TSA" and by_row[7].km == pytest.approx(256, abs=1)
    assert by_row[8].status == "filled" and by_row[8].dest_iata == "KHH"
    assert by_row[9].status == "filled" and by_row[9].km == pytest.approx(12566, abs=5)
    assert by_row[10].status == "filled" and by_row[10].dest_iata == "HKG"
    assert by_row[2].status == "unresolved" and by_row[11].status == "unresolved"
    assert by_row[12].status == "hsr" and by_row[12].km == 999
    assert by_row[13].status == "kept" and by_row[13].km == 123
    assert by_row[16].status == "filled" and by_row[16].dest_iata == "SAI"
    assert by_row[17].status == "filled" and by_row[17].km == 0
    assert by_row[18].status == "filled" and by_row[18].origin_iata == "TSA" and by_row[18].km == pytest.approx(256, abs=1)
    assert by_row[19].status == "filled"
    assert dict(summary.unresolved) == {"SHA": 1, "火星": 1}
    assert dict(summary.unresolved_known) == {}
    assert "火星,,,   # 出現 1 次" in summary.console_summary()

    ws = openpyxl.load_workbook(out)["航空"]
    assert ws["G4"].value == 132 and isinstance(ws["G4"].value, int)
    assert ws["G5"].value is None and ws["G6"].value is None           # HSR rows left blank
    assert ws["G12"].value == 999 and ws["G13"].value == 123           # existing kept
    assert ws["G3"].value is None and ws["G14"].value is None          # skipped
    assert ws["J3"].value == "排除：人次為0或起迄為「無」"             # not duplicated/changed
    assert ws["J5"].value == "疑似高鐵(新竹↔台北)，請確認交通方式並移至高鐵表"
    assert ws["J6"].value == "疑似高鐵(台中↔台北)，請確認交通方式並移至高鐵表"
    assert "TAS→TSA 疑為代碼誤植" in ws["J7"].value and "疑為 TSA 誤植" in ws["J7"].value
    assert ws["J8"].value == "既有備註"                                 # untouched (no note needed)
    assert ws["J11"].value == "無法對應機場代碼：火星，請更新 " + str(lookup.path)
    assert "疑似高鐵(TPE↔KHH)" in ws["J12"].value
    assert ws["J16"].value == "迄站 SIA→SAI 疑為代碼誤植（SIA 已停用）請確認"
    assert ws["J18"].value == "起站 台北→TSA：台北預設松山；國際線請改 TPE"
    assert "PNH" in ws["J19"].value and "無定期航班" in ws["J19"].value
    # HSR rows get ONLY the HSR note, not the 台北→TSA review note
    assert "台北→TSA" not in ws["J5"].value and "台北→TSA" not in ws["J6"].value
    # preserved content
    assert ws["I2"].value == "=距離*人次*係數" and ws["I2"].data_type == "s"
    assert ws["I3"].value == "=G3*H3*0.123" and ws["I3"].data_type == "f"
    assert ws["L2"].value == "https://www.greatcirclemap.com/" and ws["K5"].value == "誤植，交通方式應調整為高鐵"
    assert ws["A1"].font.name == "微軟正黑體" and ws["A1"].font.bold
    assert ws.column_dimensions["E"].width == pytest.approx(15.9)
    other = openpyxl.load_workbook(out)["高鐵"]
    assert [c.value for c in other[2]][:5] == ["測試公司", "台北", "左營", 345, 2] and other["F2"].value == "=D2*E2"


def test_rerun_is_idempotent(workbook, lookup, tmp_path):
    out1 = tmp_path / "pass1.xlsx"
    out2 = tmp_path / "pass2.xlsx"
    run(workbook, lookup, out1)
    s2 = run(out1, lookup, out2)
    ws1 = openpyxl.load_workbook(out1)["航空"]
    ws2 = openpyxl.load_workbook(out2)["航空"]
    for r in range(1, ws1.max_row + 1):
        for c in range(1, ws1.max_column + 1):
            assert ws1.cell(r, c).value == ws2.cell(r, c).value, ws1.cell(r, c).coordinate
    assert s2.counts["kept"] >= 1 and s2.counts["filled"] == 0  # everything already filled is kept


def test_overwrite_and_fill_hsr(workbook, lookup, tmp_path):
    out = tmp_path / "ow.xlsx"
    s = run(workbook, lookup, out, overwrite=True, fill_hsr_distance=True, method="haversine", decimals=1)
    by_row = {r.row: r for r in s.results}
    ws = openpyxl.load_workbook(out)["航空"]
    assert ws["G13"].value != 123 and ws["G13"].value == pytest.approx(255.8, abs=1.0)
    assert by_row[12].status == "hsr" and ws["G12"].value == pytest.approx(300, abs=10)   # TPE-KHH filled
    assert by_row[6].status == "hsr" and isinstance(ws["G6"].value, float)                 # RMQ-TSA filled
    assert by_row[5].status == "unresolved"   # 新竹 still has no airport even with --fill-hsr-distance
    assert dict(s.unresolved_known) == {"新竹": 1} and "需人工決定" in s.console_summary()


def test_overrides_take_precedence(workbook, lookup, tmp_path):
    ov_path = tmp_path / "ov.csv"
    ov_path.write_text("from,to,km,source\n#comment\nJFK,TPE,12570,gcm\n", encoding="utf-8-sig")
    ov = Overrides.load(ov_path)
    s = fill_workbook(workbook, tmp_path / "ov.xlsx", lookup, ov, Options())
    r = {x.row: x for x in s.results}[9]
    assert r.km == 12570 and "distance_overrides" in r.note and "gcm" in r.note


def test_output_equals_input_refused(workbook, lookup):
    with pytest.raises(ValueError):
        run(workbook, lookup, workbook)


def test_missing_sheet(workbook, lookup, tmp_path):
    with pytest.raises(ValueError):
        run(workbook, lookup, tmp_path / "x.xlsx", sheet="不存在")


def test_report_csv(workbook, lookup, tmp_path):
    s = run(workbook, lookup, tmp_path / "r.xlsx")
    rp = write_report(s, tmp_path / "r.csv")
    raw = rp.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(rp.open(encoding="utf-8-sig", newline="")))
    assert len(rows) == len(ROWS)
    assert rows[0].keys() >= {"row", "origin_raw", "dest_raw", "origin_iata", "dest_iata", "km", "status", "url", "note"}
    assert any(r["status"] == "hsr" for r in rows)
    assert "ultimately" not in s.console_summary() and "火星" in s.console_summary()


def test_default_output_path():
    assert default_output_path(Path("/x/碳盤查.xlsx")) == Path("/x/碳盤查_距離.xlsx")


def test_cli_end_to_end(workbook, lookup_path, tmp_path, monkeypatch, capsys):
    # synthetic airports.csv for the CLI
    from tests.conftest import SYNTH_AIRPORTS
    ap = tmp_path / "airports.csv"
    with ap.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["iata", "name", "municipality", "iso_country", "type", "lat", "lon", "scheduled_service"])
        for a in SYNTH_AIRPORTS:
            w.writerow([a.iata, a.name, a.municipality, a.iso_country, a.type, a.lat, a.lon, "yes" if a.scheduled_service else "no"])
    out = tmp_path / "cli_out.xlsx"
    rc = main(["fill", str(workbook), "-o", str(out), "--lookup", str(lookup_path), "--airports", str(ap)])
    assert rc == 0
    assert out.exists() and (tmp_path / "cli_out_report.csv").exists()
    text = capsys.readouterr().out
    assert "filled=" in text and "火星" in text

    # default output path next to the input, never the input itself
    rc = main(["fill", str(workbook), "--lookup", str(lookup_path), "--airports", str(ap), "--no-report"])
    assert rc == 0 and (tmp_path / "synthetic_距離.xlsx").exists()

    rc = main(["fill", str(tmp_path / "missing.xlsx"), "--lookup", str(lookup_path), "--airports", str(ap)])
    assert rc == 1 and "錯誤" in capsys.readouterr().err
