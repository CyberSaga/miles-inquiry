"""碳排量 = 距離*人次*係數 (航空 by category from K/L, 高鐵 from J2) – synthetic workbooks only."""
from __future__ import annotations

import csv

import openpyxl

from miles_inquiry import carbon
from miles_inquiry.fill import Options, fill_workbook, write_report
from miles_inquiry.overrides import Overrides

from test_fill import build_workbook


def build_with_factors(path):
    build_workbook(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["航空"]
    for r, (label, value) in enumerate((("商務差旅(飛機)-國際線", 0.123), ("商務差旅(飛機)-國內線", 0.2),
                                        ("商務差旅(飛機)-中國分行", 0.088)), start=3):
        ws.cell(r, 11).value, ws.cell(r, 12).value = label, value
    ws["I12"] = 42                                  # a value someone typed by hand: never overwritten
    hsr = wb["高鐵"]
    hsr["F1"], hsr["G1"], hsr["H1"] = "碳排量", "備註", None
    hsr["F2"] = "=距離*人次*係數"                    # template text, replaced
    hsr["G2"] = 0.029                               # factor cell (J2 in the real workbook)
    hsr.append(["測試公司", "台中", "台北", None, 1])      # 3: no distance yet → formula waits
    hsr.append(["測試公司", "無", "無", None, 0])          # 4: excluded → blank
    hsr.append(["測試公司", "無", "無", None, 1])          # 5: 無 with 人次 1 → blank
    wb.save(path)
    return path


def run(path, lookup, out, **kw):
    kw.setdefault("hsr_factor_cell", "G2")
    return fill_workbook(path, out, lookup, Overrides({}), Options(**kw))


def test_classify():
    assert carbon.classify("TW", "TW") == carbon.DOMESTIC
    assert carbon.classify("CN", "CN") == carbon.CHINA
    assert carbon.classify("CN", "TW") == carbon.INTL          # cross-strait
    assert carbon.classify("HK", "TW") == carbon.INTL
    assert carbon.classify("US", "US") == carbon.INTL          # overseas branches' domestic flights


def test_writable_and_formula():
    assert carbon.writable(None) and carbon.writable("  ")
    assert carbon.writable("=距離*人次*係數")
    assert carbon.writable("=G5*H5*$L$4") and carbon.writable('=IF(G5="","",G5*H5*$J$2)')
    assert not carbon.writable("=G3*H3*0.123") and not carbon.writable(42)
    assert carbon.writable(42, overwrite=True)
    f = carbon.Factor("$L$4", 0.2, "")
    assert carbon.formula(7, 7, 8, f) == "=G7*H7*$L$4"
    assert carbon.formula(7, 7, 8, f, blank_guard=True) == '=IF(G7="","",G7*H7*$L$4)'


def test_find_air_factors_by_label(tmp_path):
    p = build_with_factors(tmp_path / "w.xlsx")
    found = carbon.find_air_factors(openpyxl.load_workbook(p)["航空"])
    assert {k: (f.ref, f.value) for k, f in found.items()} == {
        carbon.INTL: ("$L$3", 0.123), carbon.DOMESTIC: ("$L$4", 0.2), carbon.CHINA: ("$L$5", 0.088)}


def test_carbon_end_to_end(lookup, tmp_path):
    src = build_with_factors(tmp_path / "w.xlsx")
    out = tmp_path / "o.xlsx"
    s = run(src, lookup, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["航空"]
    assert ws["I4"].value == "=G4*H4*$L$4"          # MZG-KHH 國內線
    assert ws["I9"].value == "=G9*H9*$L$3"          # TPE-JFK 國際線
    assert ws["I10"].value == "=G10*H10*$L$3"       # TPE-HKG 國際線
    assert ws["I13"].value == "=G13*H13*$L$4"       # kept distance still gets a formula
    assert ws["I2"].value == "=距離*人次*係數"       # unresolved row: template left as is
    assert ws["I3"].value == "=G3*H3*0.123"         # someone's own formula on a skipped row
    assert ws["I5"].value is None and ws["I11"].value is None and ws["I14"].value is None   # HSR / unresolved / skipped
    assert ws["I12"].value == 42 and s.carbon_kept == []   # row 12 is HSR-like anyway
    assert wb.calculation.fullCalcOnLoad

    by_row = {r.row: r for r in s.results}
    assert (by_row[4].category, by_row[4].factor, by_row[4].carbon) == ("國內線", 0.2, round(132 * 2 * 0.2, 6))
    assert (by_row[9].category, by_row[9].carbon) == ("國際線", round(12566 * 3 * 0.123, 6))
    assert s.carbon_counts["國內線"] >= 1 and s.carbon_counts["國際線"] >= 1

    hsr = wb["高鐵"]
    assert hsr["F2"].value == '=IF(D2="","",D2*E2*$G$2)'
    assert hsr["F3"].value == '=IF(D3="","",D3*E3*$G$2)'
    assert hsr["F4"].value is None and hsr["F5"].value is None
    assert s.hsr_carbon_rows == 2 and s.hsr_carbon_pending == 1
    assert "碳排量（高鐵）：2 列，係數 0.029" in s.console_summary()

    rows = list(csv.DictReader(write_report(s, tmp_path / "r.csv").read_text(encoding="utf-8-sig").splitlines()))
    assert {"category", "factor", "carbon"} <= rows[0].keys()
    assert next(r for r in rows if r["row"] == "4")["carbon"] == str(round(132 * 2 * 0.2, 6))


def test_user_value_kept_unless_overwrite(lookup, tmp_path):
    src = build_with_factors(tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(src)
    wb["航空"]["I9"] = 999
    wb.save(src)
    s = run(src, lookup, tmp_path / "a.xlsx")
    assert openpyxl.load_workbook(tmp_path / "a.xlsx")["航空"]["I9"].value == 999 and s.carbon_kept == [9]
    run(src, lookup, tmp_path / "b.xlsx", overwrite=True)
    assert openpyxl.load_workbook(tmp_path / "b.xlsx")["航空"]["I9"].value == "=G9*H9*$L$3"


def test_rerun_keeps_formulas_and_no_carbon_flag(lookup, tmp_path):
    src = build_with_factors(tmp_path / "w.xlsx")
    run(src, lookup, tmp_path / "1.xlsx")
    s2 = run(tmp_path / "1.xlsx", lookup, tmp_path / "2.xlsx")
    a, b = (openpyxl.load_workbook(tmp_path / n) for n in ("1.xlsx", "2.xlsx"))
    for name in ("航空", "高鐵"):
        assert [[c.value for c in r] for r in a[name].iter_rows()] == [[c.value for c in r] for r in b[name].iter_rows()]
    assert s2.carbon_kept == []
    run(src, lookup, tmp_path / "n.xlsx", carbon=False)
    ws = openpyxl.load_workbook(tmp_path / "n.xlsx")
    assert ws["航空"]["I4"].value is None and ws["高鐵"]["F2"].value == "=距離*人次*係數"


def test_missing_factor_or_sheet_warns(lookup, tmp_path):
    src = build_workbook(tmp_path / "plain.xlsx")      # no factor table, 高鐵 has no 碳排量 column
    s = run(src, lookup, tmp_path / "o.xlsx", hsr_factor_cell="J2")
    assert any("找不到 國際線、國內線、中國分行 的碳排係數" in w for w in s.warnings)
    assert any("高鐵" in w and "碳排量" in w for w in s.warnings)
    assert openpyxl.load_workbook(tmp_path / "o.xlsx")["航空"]["I4"].value is None
