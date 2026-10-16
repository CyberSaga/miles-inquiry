"""Core: open the workbook, resolve codes, compute distances, write notes/report."""
from __future__ import annotations

import csv
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from . import carbon
from .geo import DEFAULT_METHOD, METHODS, distance_km, round_km
from .hsr import hsr_note, is_hsr_like
from .lookup import Lookup, Resolution, is_skip_value
from .overrides import Overrides
from .rail import LOOKUP_FILENAME as RAIL_LOOKUP_FILENAME
from .rail import OVERRIDES_FILENAME as RAIL_OVERRIDES_FILENAME
from .rail import RailData

GCM_URL = "https://www.greatcirclemap.com/?routes={o}-{d}"
NOTE_SEP = "；"
DEFAULT_SHEET = "航空"
DEFAULT_HSR_SHEET = "高鐵"
DEFAULT_HSR_FACTOR_CELL = "J2"
SUPPORTED_SUFFIXES = (".xlsx", ".xlsm")

# header text (normalised: NFKC, whitespace removed) → role
HEADER_ALIASES: Dict[str, Tuple[str, ...]] = {
    "origin": ("航班(起站)", "起站", "航班起站"),
    "dest": ("航班(迄站)", "迄站", "航班迄站"),
    "distance": ("距離", "距離(km)", "距離km"),
    "pax": ("人次",),
    "note": ("備註",),
    "carbon": ("碳排量", "碳排放量"),
}
HSR_HEADER_ALIASES: Dict[str, Tuple[str, ...]] = {
    "origin": ("高鐵(起站)", "起站"),
    "dest": ("高鐵(迄站)", "迄站"),
    "distance": HEADER_ALIASES["distance"],
    "pax": HEADER_ALIASES["pax"],
    "carbon": HEADER_ALIASES["carbon"],
    "note": HEADER_ALIASES["note"],
}
REQUIRED_ROLES = ("origin", "dest", "distance")

STATUSES = ("filled", "kept", "hsr", "unresolved", "skipped")


def _norm_header(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", "", text)


@dataclass
class Options:
    sheet: str = DEFAULT_SHEET
    method: str = DEFAULT_METHOD
    decimals: int = 0
    overwrite: bool = False
    fill_hsr_distance: bool = False
    convert_names: bool = True     # replace 澎湖 / HONG KONG / TAS … in the station cells by the IATA code
    carbon: bool = True            # 碳排量 = 距離*人次*係數 (Excel formulas) on 航空 and 高鐵
    hsr_sheet: str = DEFAULT_HSR_SHEET
    hsr_factor_cell: str = DEFAULT_HSR_FACTOR_CELL
    hsr_distance: bool = True      # fill 高鐵 距離 from data/thsr_distances.csv + rail_distance_overrides.csv
    max_header_scan: int = 20

    def __post_init__(self):
        if self.method not in METHODS:
            raise ValueError(f"method 必須是 {METHODS} 之一，收到 {self.method!r}")


@dataclass
class RowResult:
    row: int
    origin_raw: str
    dest_raw: str
    origin_iata: Optional[str]
    dest_iata: Optional[str]
    km: Optional[float]
    status: str
    url: str = ""
    note: str = ""          # everything worth telling (report)
    cell_notes: List[str] = field(default_factory=list)   # what went into 備註
    category: str = ""      # 國際線 / 國內線 / 中國分行 (碳排係數類別)
    factor: Optional[float] = None
    carbon: Optional[float] = None   # 距離*人次*係數, as Excel will compute it

    def as_report_row(self) -> dict:
        return {
            "row": self.row,
            "origin_raw": self.origin_raw,
            "dest_raw": self.dest_raw,
            "origin_iata": self.origin_iata or "",
            "dest_iata": self.dest_iata or "",
            "km": "" if self.km is None else self.km,
            "status": self.status,
            "url": self.url,
            "note": self.note,
            "category": self.category,
            "factor": "" if self.factor is None else self.factor,
            "carbon": "" if self.carbon is None else self.carbon,
        }


REPORT_COLUMNS = ("row", "origin_raw", "dest_raw", "origin_iata", "dest_iata", "km", "status", "url", "note",
                  "category", "factor", "carbon")


@dataclass
class FillSummary:
    input_path: Path
    output_path: Path
    sheet: str
    header_row: int
    columns: Dict[str, int]
    results: List[RowResult]
    counts: Counter = field(default_factory=Counter)
    unresolved: Counter = field(default_factory=Counter)   # raw value → occurrences (not in lookup.csv)
    unresolved_known: Counter = field(default_factory=Counter)   # raw value → occurrences (lookup.csv row with empty iata)
    blank_rows: List[int] = field(default_factory=list)          # rows with an empty 起站/迄站 next to a real trip
    warnings: List[str] = field(default_factory=list)
    carbon_counts: Counter = field(default_factory=Counter)   # 航空 category label → rows with a 碳排量 formula
    carbon_kept: List[int] = field(default_factory=list)       # rows whose 碳排量 already held other content
    hsr_carbon_rows: int = 0
    hsr_carbon_pending: int = 0                                # 高鐵 rows whose 距離 is still blank
    hsr_factor: Optional[carbon.Factor] = None
    hsr_distance_counts: Counter = field(default_factory=Counter)   # thsr / override / kept / missing
    hsr_missing_pairs: Dict[Tuple[str, str], List[int]] = field(default_factory=dict)   # pair → rows
    hsr_unresolved: Counter = field(default_factory=Counter)       # value → occurrences (not a station)
    hsr_missing_thsr: Counter = field(default_factory=Counter)       # THSR pair missing from thsr_distances.csv

    def console_summary(self) -> str:
        lines = [
            f"輸入：{self.input_path}",
            f"輸出：{self.output_path}",
            f"工作表「{self.sheet}」標題列 {self.header_row}，共處理 {len(self.results)} 列",
            "  " + "、".join(f"{s}={self.counts.get(s, 0)}" for s in STATUSES),
        ]
        if self.carbon_counts:
            lines.append("碳排量（航空）：" + "、".join(f"{k} {v} 列" for k, v in self.carbon_counts.items()))
        if self.carbon_kept:
            lines.append("碳排量欄已有其他內容、保留未改的列：" + "、".join(str(r) for r in self.carbon_kept))
        if self.hsr_distance_counts:
            c = self.hsr_distance_counts
            lines.append(f"距離（高鐵）：台灣高鐵站間距離表 {c['thsr']} 列、{RAIL_OVERRIDES_FILENAME} {c['override']} 列、"
                         f"保留既有 {c['kept']} 列、無法計算 {c['missing']} 列")
        if self.hsr_missing_thsr:
            lines.append("data/thsr_distances.csv 缺少的站對（相關列距離留空）：" +
                         "、".join(f"{k}（{v} 列）" for k, v in self.hsr_missing_thsr.most_common()))
        if self.hsr_unresolved:
            lines.append(f"高鐵表無法對應的站名（請加入 data/{RAIL_LOOKUP_FILENAME}）：")
            for raw, n in self.hsr_unresolved.most_common():
                lines.append(f"  {raw}   # 出現 {n} 次")
        if self.hsr_missing_pairs:
            lines.append(f"高鐵表缺距離的站對 {len(self.hsr_missing_pairs)} 組（請填 data/{RAIL_OVERRIDES_FILENAME}）")
        if self.hsr_factor is not None:
            line = f"碳排量（高鐵）：{self.hsr_carbon_rows} 列，係數 {self.hsr_factor.value}（{self.hsr_factor.label}）"
            if self.hsr_carbon_pending:
                line += f"；其中 {self.hsr_carbon_pending} 列距離空白，填入距離後公式自動計算"
            lines.append(line)
        for w in self.warnings:
            lines.append(f"警告：{w}")
        if self.unresolved:
            lines.append("無法對應的值（請加入 data/lookup.csv，格式 value,iata,review,note）：")
            for raw, n in self.unresolved.most_common():
                lines.append(f"  {raw},,,   # 出現 {n} 次")
        if self.unresolved_known:
            lines.append("lookup.csv 已標為無法對應、需人工決定代碼的值：")
            for raw, n in self.unresolved_known.most_common():
                lines.append(f"  {raw}   # 出現 {n} 次")
        if self.blank_rows:
            lines.append("起站或迄站空白、請補填的列：" + "、".join(str(r) for r in self.blank_rows))
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def find_header(ws: Worksheet, max_scan: int = 20,
                aliases_by_role: Optional[Dict[str, Tuple[str, ...]]] = None) -> Tuple[int, Dict[str, int]]:
    """Return (header_row_index, {role: column_index}) – 1-based."""
    alias_to_role = {}
    for role, aliases in (aliases_by_role or HEADER_ALIASES).items():
        for a in aliases:
            alias_to_role[_norm_header(a)] = role
    for r in range(1, min(ws.max_row, max_scan) + 1):
        found: Dict[str, int] = {}
        for c in range(1, ws.max_column + 1):
            role = alias_to_role.get(_norm_header(ws.cell(r, c).value))
            if role and role not in found:
                found[role] = c
        if all(role in found for role in REQUIRED_ROLES):
            return r, found
    raise ValueError(
        f"在工作表「{ws.title}」前 {max_scan} 列找不到標題（需含 起站、迄站、距離）"
    )


def split_notes(text) -> List[str]:
    if text is None:
        return []
    parts = re.split(r"[；;]", str(text))
    return [p.strip() for p in parts if p.strip()]


def _contains_run(parts: List[str], run: List[str]) -> bool:
    """True when *run* appears as a contiguous sub-sequence of *parts*."""
    n = len(run)
    return n > 0 and any(parts[i:i + n] == run for i in range(len(parts) - n + 1))


def append_note(existing, note: str) -> str:
    """Append *note* to an existing 備註 value, without duplicating.

    The existing text is kept exactly as written (its own separators, spacing
    and wording are never rewritten); ``split_notes`` is used only to detect
    that the note is already there.  A note may itself contain ``；`` (e.g. a
    multi-sentence lookup note), so the duplicate check compares segment
    sequences rather than whole strings.
    """
    existing_text = "" if existing is None else str(existing)
    new_parts = split_notes(note)
    if not new_parts or _contains_run(split_notes(existing_text), new_parts):
        return existing_text
    note = note.strip()
    if not existing_text.strip():
        return note
    if re.search(r"[；;]\s*$", existing_text):
        return existing_text + note
    return existing_text + NOTE_SEP + note


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _pax_value(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = unicodedata.normalize("NFKC", str(value)).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def airport_source_note(ap) -> str:
    """備註 text when the coordinates do not come from greatcirclemap.com's own data."""
    if ap is None or ap.on_site:
        return ""
    if ap.conflict_note:
        return (f"{ap.iata} 代碼已重新指派：{ap.conflict_note}，座標取自 OurAirports（{ap.name}），"
                f"greatcirclemap 網站數字不適用")
    return f"{ap.iata} 不在 greatcirclemap 資料庫，座標取自 OurAirports（{ap.name}）"


def converted_note(role: str, original: str, iata: str) -> str:
    """備註 text when a station cell was rewritten from a name to its IATA code."""
    return f"{role}原值「{original}」→{iata}"


def unresolved_note(res: Resolution, lookup_name: str = "data/lookup.csv", role: str = "") -> str:
    if res.status == "blank":
        return f"{role or '起站/迄站'}空白，請補填"
    if res.via == "lookup":  # explicitly marked unresolvable in lookup.csv
        return f"無法對應機場代碼：{res.raw}" + (f"（{res.note}）" if res.note else "")
    return f"無法對應機場代碼：{res.raw}，請更新 {lookup_name}"


def lookup_display_name(path) -> str:
    """Short name for 備註 (``data/lookup.csv``) – never a full machine path."""
    if not path:
        return "data/lookup.csv"
    path = Path(path)
    return f"{path.parent.name}/{path.name}" if path.parent.name else path.name


# ---------------------------------------------------------------------------
# main entry
# ---------------------------------------------------------------------------

def default_output_path(input_path: Path) -> Path:
    input_path = Path(input_path)
    return input_path.with_name(f"{input_path.stem}_距離{input_path.suffix or '.xlsx'}")


def default_report_path(output_path: Path) -> Path:
    output_path = Path(output_path)
    return output_path.with_name(f"{output_path.stem}_report.csv")


def next_free_path(path: Path) -> Path:
    """``x.xlsx`` → ``x.xlsx`` if free, else ``x(2).xlsx``, ``x(3).xlsx`` … (never clobbers an existing file)."""
    path = Path(path)
    if not path.exists():
        return path
    stem = re.sub(r"\(\d+\)$", "", path.stem)
    n = 2
    while True:
        candidate = path.with_name(f"{stem}({n}){path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def load_input_workbook(input_path: Path):
    """openpyxl load with Chinese errors; .xlsm keeps its VBA project."""
    suffix = input_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"{input_path.name}：只支援 {' / '.join(SUPPORTED_SUFFIXES)}（不支援 .xls、.csv）；"
                         "請先在 Excel 另存為 .xlsx")
    try:
        # NOT data_only → formulas preserved; keep_vba so an .xlsm stays a valid macro workbook
        return openpyxl.load_workbook(input_path, keep_vba=(suffix == ".xlsm"))
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise ValueError(f"{input_path} 不是有效的 {suffix} 檔（可能是 .xls/.csv 改了副檔名，或檔案損毀）：{exc}") from exc


def fill_workbook(input_path: Path, output_path: Path, lookup: Lookup,
                  overrides: Optional[Overrides] = None,
                  options: Optional[Options] = None,
                  rail: Optional[RailData] = None) -> FillSummary:
    options = options or Options()
    overrides = overrides or Overrides({})
    input_path, output_path = Path(input_path), Path(output_path)
    if not input_path.is_file():
        raise FileNotFoundError(f"找不到輸入檔 {input_path}")
    if output_path.resolve() == input_path.resolve():
        raise ValueError("輸出檔不可與輸入檔相同（工具絕不覆寫原始檔）")
    if input_path.suffix.lower() == ".xlsm" and output_path.suffix.lower() != ".xlsm":
        raise ValueError("輸入為 .xlsm（含巨集）時，輸出檔也必須是 .xlsm，否則 Excel 無法開啟")

    wb = load_input_workbook(input_path)
    if options.sheet not in wb.sheetnames:
        raise ValueError(f"活頁簿沒有工作表「{options.sheet}」，現有：{wb.sheetnames}")
    ws = wb[options.sheet]
    header_row, cols = find_header(ws, options.max_header_scan)

    summary = FillSummary(input_path, output_path, options.sheet, header_row, cols, [])
    if "note" not in cols:
        summary.warnings.append("找不到「備註」欄，說明只會寫入報表，不會寫入活頁簿")
    if "pax" not in cols:
        summary.warnings.append("找不到「人次」欄，無法依人次=0 排除")

    lookup_name = lookup_display_name(lookup.path)

    for r in range(header_row + 1, ws.max_row + 1):
        o_cell = ws.cell(r, cols["origin"])
        d_cell = ws.cell(r, cols["dest"])
        dist_cell = ws.cell(r, cols["distance"])
        pax = _pax_value(ws.cell(r, cols["pax"]).value) if "pax" in cols else None
        note_cell = ws.cell(r, cols["note"]) if "note" in cols else None

        if _is_blank(o_cell.value) and _is_blank(d_cell.value) and pax is None:
            continue  # completely empty row

        res_o = lookup.resolve(o_cell.value)
        res_d = lookup.resolve(d_cell.value)
        result = RowResult(r, res_o.raw, res_d.raw, res_o.iata, res_d.iata, None, "")
        report_notes: List[str] = []

        def add_cell_note(text: str):
            result.cell_notes.append(text)
            if note_cell is not None:
                note_cell.value = append_note(note_cell.value, text)

        # 1. skipped rows ---------------------------------------------------
        if res_o.status == "skip" or res_d.status == "skip" or (pax is not None and pax == 0):
            result.status = "skipped"
            report_notes.append("排除：人次為0或起迄為「無」")
            _finish(result, report_notes, summary)
            continue

        if res_o.iata and res_d.iata:
            result.url = GCM_URL.format(o=res_o.iata, d=res_d.iata)

        # 1b. blank endpoint next to a real trip → data-entry gap, ask a human ----
        blanks = [(role, res) for role, res in (("起站", res_o), ("迄站", res_d)) if res.status == "blank"]
        if blanks:
            for role, res in blanks:
                add_cell_note(unresolved_note(res, lookup_name, role))
            summary.blank_rows.append(r)
            result.status = "unresolved"
            if not _is_blank(dist_cell.value):
                result.km = dist_cell.value
                report_notes.append("已有距離，保留")
            _finish(result, report_notes, summary)
            continue

        # 2. HSR-like (checked first: an HSR row gets only the HSR note) ------
        hsr = is_hsr_like(o_cell.value, d_cell.value, res_o.iata, res_d.iata)
        if hsr:
            add_cell_note(hsr_note(o_cell.value, d_cell.value))
            result.status = "hsr"
            if not options.fill_hsr_distance:
                if not _is_blank(dist_cell.value):
                    result.km = dist_cell.value
                    report_notes.append("已有距離，保留")
                _finish(result, report_notes, summary)
                continue

        # 3. mapping / correction / review notes ----------------------------
        for role, res, cell in (("起站", res_o, o_cell), ("迄站", res_d, d_cell)):
            if res.status != "ok":
                continue
            cell_note = res.note_for_cell
            # 3a. a name or mistyped code becomes the IATA code; the original value is kept in 備註.
            # HSR rows keep their text (they are train trips, and 新竹 has no airport at all).
            if options.convert_names and not hsr and res.mapped:
                original = str(cell.value).strip()
                cell.value = res.iata
                if not cell_note:   # correction / review notes already read "原值→代碼"
                    add_cell_note(converted_note(role, original, res.iata))
            if cell_note:
                add_cell_note(f"{role} {cell_note}")
            else:
                report_note = res.note_for_report
                if report_note:
                    report_notes.append(f"{role} {report_note}")
            if res.via == "direct":
                ap = lookup.airports.get(res.iata)
                if ap is not None and not ap.scheduled_service:
                    add_cell_note(f"{role} {res.iata}（{ap.name}, {ap.iso_country}）無定期航班，請確認代碼")

        # 4. unresolved -----------------------------------------------------
        unresolved = [res for res in (res_o, res_d) if res.status == "unresolved"]
        if unresolved:
            for res in unresolved:
                add_cell_note(unresolved_note(res, lookup_name))
                (summary.unresolved_known if res.via == "lookup" else summary.unresolved)[res.raw] += 1
            result.status = "unresolved"
            if not _is_blank(dist_cell.value):
                result.km = dist_cell.value
                report_notes.append("已有距離，保留")
            _finish(result, report_notes, summary)
            continue

        # 5. distance -------------------------------------------------------
        if not _is_blank(dist_cell.value) and not options.overwrite:
            result.km = dist_cell.value
            if result.status != "hsr":
                result.status = "kept"
            report_notes.append("已有距離，保留（--overwrite 可覆寫）")
            _finish(result, report_notes, summary)
            continue

        ov = overrides.get(res_o.iata, res_d.iata)
        if ov is not None:
            km = round_km(ov.km, options.decimals)
            report_notes.append("距離來自 distance_overrides.csv" + (f"（{ov.source}）" if ov.source else ""))
        else:
            a = lookup.airports.get(res_o.iata)
            b = lookup.airports.get(res_d.iata)
            for role, ap in (("起站", a), ("迄站", b)):
                src_note = airport_source_note(ap)
                if src_note:
                    add_cell_note(f"{role} {src_note}")
            km = round_km(distance_km(a.lat, a.lon, b.lat, b.lon, options.method), options.decimals)
            if res_o.iata == res_d.iata:
                report_notes.append("起迄相同，距離為 0")
        dist_cell.value = km
        result.km = km
        if result.status != "hsr":
            result.status = "filled"
        _finish(result, report_notes, summary)

    if options.hsr_distance and rail is not None:
        _fill_hsr_distance(wb, rail, summary, options)
    if options.carbon:
        _fill_air_carbon(ws, cols, lookup, summary, options)
        _fill_hsr_carbon(wb, summary, options)
        if summary.carbon_counts or summary.hsr_carbon_rows:
            wb.calculation.fullCalcOnLoad = True   # no cached values from openpyxl → Excel computes on open
    wb.save(output_path)
    return summary


def _fill_air_carbon(ws: Worksheet, cols: Dict[str, int], lookup: Lookup, summary: FillSummary,
                     options: Options) -> None:
    if "carbon" not in cols or "pax" not in cols:
        summary.warnings.append(f"「{ws.title}」找不到「碳排量」或「人次」欄，未填碳排量")
        return
    factors = carbon.find_air_factors(ws)
    missing = [carbon.CATEGORY_LABELS[c] for c in carbon.CATEGORY_LABELS if c not in factors]
    if missing:
        summary.warnings.append(f"「{ws.title}」找不到 {'、'.join(missing)} 的碳排係數（類別文字右側需為數字），"
                                "這些列的碳排量留空")
    for res in summary.results:
        if res.status not in ("filled", "kept") or not isinstance(res.km, (int, float)) or isinstance(res.km, bool):
            continue
        a, b = lookup.airports.get(res.origin_iata), lookup.airports.get(res.dest_iata)
        if a is None or b is None:
            continue
        cat = carbon.classify(a.iso_country, b.iso_country)
        res.category = carbon.CATEGORY_LABELS[cat]
        factor = factors.get(cat)
        if factor is None:
            continue
        cell = ws.cell(res.row, cols["carbon"])
        if not carbon.writable(cell.value, options.overwrite):
            summary.carbon_kept.append(res.row)
            continue
        pax = _pax_value(ws.cell(res.row, cols["pax"]).value)
        cell.value = carbon.formula(res.row, cols["distance"], cols["pax"], factor)
        res.factor = factor.value
        res.carbon = None if pax is None else round(res.km * pax * factor.value, 6)
        summary.carbon_counts[res.category] += 1


def _hsr_sheet(wb, summary: FillSummary, options: Options, what: str):
    if options.hsr_sheet not in wb.sheetnames:
        summary.warnings.append(f"沒有工作表「{options.hsr_sheet}」，未填高鐵{what}")
        return None, 0, {}
    ws = wb[options.hsr_sheet]
    try:
        header_row, cols = find_header(ws, options.max_header_scan, HSR_HEADER_ALIASES)
    except ValueError as exc:
        summary.warnings.append(f"{exc}，未填高鐵{what}")
        return None, 0, {}
    return ws, header_row, cols


def _fill_hsr_distance(wb, rail: RailData, summary: FillSummary, options: Options) -> None:
    ws, header_row, cols = _hsr_sheet(wb, summary, options, "距離")
    if ws is None:
        return
    factor_coord = options.hsr_factor_cell.replace("$", "").upper()
    for r in range(header_row + 1, ws.max_row + 1):
        o, d = ws.cell(r, cols["origin"]).value, ws.cell(r, cols["dest"]).value
        pax = _pax_value(ws.cell(r, cols["pax"]).value) if "pax" in cols else None
        if _is_blank(o) and _is_blank(d):
            continue
        if is_skip_value(o) or is_skip_value(d) or (pax is not None and pax == 0):
            continue
        note_cell = ws.cell(r, cols["note"]) if "note" in cols else None
        if note_cell is not None and note_cell.coordinate == factor_coord:
            note_cell = None    # J2 holds the 碳排係數 – never write text into it

        def add_note(text: str):
            if note_cell is not None:
                note_cell.value = append_note(note_cell.value, text)

        blanks = [role for role, v in (("起站", o), ("迄站", d)) if _is_blank(v)]
        if blanks:
            for role in blanks:
                add_note(f"{role}空白，請補填")
            summary.hsr_distance_counts["missing"] += 1
            continue
        dist_cell = ws.cell(r, cols["distance"])
        if not _is_blank(dist_cell.value) and not options.overwrite:
            summary.hsr_distance_counts["kept"] += 1
            continue

        dist, ra, rb = rail.distance(o, d)
        for role, res in (("起站", ra), ("迄站", rb)):
            if res.review and not res.non_station:
                add_note(f"{role} {res.raw}→{res.canonical}" + (f"：{res.note}" if res.note else "，請確認"))
        if dist.km is not None:
            if dist.note == "起迄相同":
                add_note("起迄同站，距離為 0，請確認")
            dist_cell.value = int(dist.km) if float(dist.km).is_integer() else dist.km
            summary.hsr_distance_counts[dist.source] += 1
            continue

        summary.hsr_distance_counts["missing"] += 1
        if dist.source == "missing_pair":   # both THSR stations but the pair table lacks them
            summary.hsr_missing_thsr[dist.note] += 1
            add_note(f"data/thsr_distances.csv 缺 {dist.note} 距離")
            continue
        unknown = [(role, res) for role, res in (("起站", ra), ("迄站", rb)) if not res.station]
        non_station = [(role, res) for role, res in unknown if res.non_station]
        if non_station:
            for role, res in non_station:
                add_note(f"{role}「{res.raw}」非高鐵站" + (f"（{res.note}）" if res.note else "") + "，請確認")
            continue
        if len(unknown) == 2:               # neither end is a THSR station → e.g. China HSR
            pair = tuple(sorted((ra.canonical, rb.canonical)))
            summary.hsr_missing_pairs.setdefault(pair, []).append(r)
            add_note(f"非台灣高鐵路段（{ra.raw}→{rb.raw}），請在 data/{RAIL_OVERRIDES_FILENAME} 填距離")
            continue
        for role, res in unknown:
            summary.hsr_unresolved[res.raw] += 1
            add_note(f"無法對應高鐵站：{res.raw}，請更新 data/{RAIL_LOOKUP_FILENAME}")


def write_rail_todo(summary: FillSummary, path: Path) -> Optional[Path]:
    """Template of the station pairs that still need a distance (``from,to,km,source,rows``)."""
    if not summary.hsr_missing_pairs:
        return None
    path = Path(path)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["from", "to", "km", "source", "rows"])
        for (a, b), rows in sorted(summary.hsr_missing_pairs.items(), key=lambda kv: (-len(kv[1]), kv[0])):
            w.writerow([a, b, "", "", " ".join(str(r) for r in rows)])
    return path


def _fill_hsr_carbon(wb, summary: FillSummary, options: Options) -> None:
    ws, header_row, cols = _hsr_sheet(wb, summary, options, "碳排量")
    if ws is None:
        return
    if "carbon" not in cols or "pax" not in cols:
        summary.warnings.append(f"「{ws.title}」找不到「碳排量」或「人次」欄，未填碳排量")
        return
    factor = carbon.read_factor(ws, options.hsr_factor_cell)
    if factor is None:
        summary.warnings.append(f"「{ws.title}」{options.hsr_factor_cell} 不是數字，無法作為高鐵碳排係數，未填碳排量")
        return
    summary.hsr_factor = factor
    for r in range(header_row + 1, ws.max_row + 1):
        o, d = ws.cell(r, cols["origin"]).value, ws.cell(r, cols["dest"]).value
        pax = _pax_value(ws.cell(r, cols["pax"]).value)
        if not pax or is_skip_value(o) or is_skip_value(d) or _is_blank(o) or _is_blank(d):
            continue   # empty, excluded (人次 0 / 無) or incomplete rows stay blank
        cell = ws.cell(r, cols["carbon"])
        if not carbon.writable(cell.value, options.overwrite):
            continue
        cell.value = carbon.formula(r, cols["distance"], cols["pax"], factor, blank_guard=True)
        summary.hsr_carbon_rows += 1
        if _is_blank(ws.cell(r, cols["distance"]).value):
            summary.hsr_carbon_pending += 1


def _finish(result: RowResult, report_notes: List[str], summary: FillSummary) -> None:
    notes = list(result.cell_notes) + [n for n in report_notes if n not in result.cell_notes]
    result.note = NOTE_SEP.join(notes)
    summary.results.append(result)
    summary.counts[result.status] += 1


def write_report(summary: FillSummary, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(REPORT_COLUMNS), lineterminator="\n")
        writer.writeheader()
        for res in summary.results:
            writer.writerow(res.as_report_row())
    return path
