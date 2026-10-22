"""碳排量 = 距離 × 人次 × 係數, written as live Excel formulas.

航空: the factors sit next to their category labels (K3:L5 in the real workbook)::

    K3 商務差旅(飛機)-國際線    L3 0.123
    K4 商務差旅(飛機)-國內線    L4 0.2
    K5 商務差旅(飛機)-中國分行  L5 0.088

They are located by label, not by fixed address, so moving the table keeps working.
Each route is classified by the airports' countries: both in Taiwan → 國內線, both
in mainland China → 中國分行, anything else (HK/MO included) → 國際線.

高鐵: a single factor cell (J2 by default).  Its 距離 column may still be empty,
so the formula there is ``=IF(G2="","",G2*H2*$J$2)`` and fills itself in once a
distance is typed.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Dict, Optional

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

INTL, DOMESTIC, CHINA = "intl", "domestic", "china"
CATEGORY_LABELS = {INTL: "國際線", DOMESTIC: "國內線", CHINA: "中國分行"}
# keyword searched in the label cell → category (first match wins per category)
CATEGORY_KEYWORDS = ((DOMESTIC, "國內線"), (INTL, "國際線"), (CHINA, "中國"))

DOMESTIC_COUNTRIES = {"TW"}
CHINA_COUNTRIES = {"CN"}

_CJK_RE = re.compile(r"[㐀-鿿]")
# formulas this module writes (so a re-run may refresh them)
_OURS_RE = re.compile(r'^=(IF\([A-Z]+\d+="","",)?[A-Z]+\d+\*[A-Z]+\d+\*\$[A-Z]+\$\d+\)?$')


@dataclass(frozen=True)
class Factor:
    ref: str          # absolute reference, e.g. $L$3
    value: float
    label: str        # text of the label cell (or the cell address)


def _number(value) -> Optional[float]:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def absolute(coord: str) -> str:
    col, row = coordinate_from_string(coord.replace("$", "").upper())
    return f"${col}${row}"


def check_coord(coord, label: str) -> str:
    """``'$l$2'`` → ``'L2'``; ValueError (not a traceback) for anything that is not one cell address."""
    try:
        col, row = coordinate_from_string(str(coord).replace("$", "").strip().upper())
        column_index_from_string(col)
    except Exception:
        row = 0
    if row < 1:
        raise ValueError(f"{label}「{coord}」不是有效的儲存格位址（例如 J2）")
    return f"{col}{row}"


def find_air_factors(ws: Worksheet) -> Dict[str, Factor]:
    """Locate 國際線 / 國內線 / 中國分行 factors: a label cell with a number right of it."""
    found: Dict[str, Factor] = {}
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            for cat, kw in CATEGORY_KEYWORDS:
                if cat in found or kw not in cell.value:
                    continue
                right = ws.cell(cell.row, cell.column + 1)
                value = _number(right.value)
                if value is not None:
                    found[cat] = Factor(absolute(right.coordinate), value, cell.value.strip())
                    break
    return found


def read_factor(ws: Worksheet, coord: str) -> Optional[Factor]:
    value = _number(ws[coord.replace("$", "")].value)
    return None if value is None else Factor(absolute(coord), value, f"{ws.title}!{absolute(coord)}")


def classify(country_a: str, country_b: str) -> str:
    pair = {country_a or "", country_b or ""}
    if pair <= DOMESTIC_COUNTRIES:
        return DOMESTIC
    if pair <= CHINA_COUNTRIES:
        return CHINA
    return INTL


def writable(value, overwrite: bool = False) -> bool:
    """Blank, the template text ``=距離*人次*係數``, or a formula we wrote earlier."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    if overwrite:
        return True
    if isinstance(value, str) and value.startswith("="):
        return bool(_CJK_RE.search(value)) or bool(_OURS_RE.match(value.replace(" ", "")))
    return False


def formula(row: int, dist_col: int, pax_col: int, factor: Factor, blank_guard: bool = False) -> str:
    g = f"{get_column_letter(dist_col)}{row}"
    h = f"{get_column_letter(pax_col)}{row}"
    body = f"{g}*{h}*{factor.ref}"
    return f'=IF({g}="","",{body})' if blank_guard else f"={body}"
